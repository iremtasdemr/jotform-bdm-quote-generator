from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_DIR = Path("output")
OUTPUT_PATH = OUTPUT_DIR / "Jotform_Enterprise_Quote_Generator_Guide.docx"

FONT_NAME = "Arial"
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x55, 0x55, 0x55)


def set_run_font(run, size=None, bold=None, color=BLACK):
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_spacing(paragraph, before=0, after=8, line=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_numbering_definition(doc, abstract_id, num_id, ordered):
    numbering = doc.part.numbering_part.element

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")

    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)

    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    level.append(num_fmt)

    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1." if ordered else "●")
    level.append(level_text)

    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)

    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "276")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    level.append(ppr)

    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_NAME)
    fonts.set(qn("w:hAnsi"), FONT_NAME)
    rpr.append(fonts)
    level.append(rpr)

    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)


def set_list_numbering(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is not None:
        ppr.remove(num_pr)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    ppr.append(num_pr)


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph)
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(doc, text, num_id=91):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, after=4)
    set_list_numbering(paragraph, num_id)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_step(doc, title, explanation, num_id=92):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, after=6)
    set_list_numbering(paragraph, num_id)
    title_run = paragraph.add_run(title)
    set_run_font(title_run, bold=True)
    body_run = paragraph.add_run(f" {explanation}")
    set_run_font(body_run)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_note(doc, label, text):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, before=4, after=10)
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.25)
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run)
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "DADCE0")
    borders.append(left)
    ppr.append(borders)
    return paragraph


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    heading_tokens = {
        1: (20, 20, 6, BLACK),
        2: (16, 18, 6, BLACK),
        3: (14, 16, 4, RGBColor(0x43, 0x43, 0x43)),
    }
    for level, (size, before, after, color) in heading_tokens.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True


def build_document():
    OUTPUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_styles(doc)
    add_numbering_definition(doc, 91, 91, ordered=False)
    for ordered_num_id in (92, 93, 94, 95):
        add_numbering_definition(
            doc,
            ordered_num_id,
            ordered_num_id,
            ordered=True,
        )

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    title.paragraph_format.line_spacing = 1.0
    title_run = title.add_run("Jotform Enterprise Quote Generator")
    set_run_font(title_run, size=26, bold=False)

    subtitle = doc.add_paragraph()
    set_paragraph_spacing(subtitle, after=18)
    subtitle_run = subtitle.add_run("A simple guide for creating clear, customer-ready quotes")
    set_run_font(subtitle_run, size=11, color=GRAY)

    intro = add_body(
        doc,
        "This guide explains what the Quote Generator is for, how to use it from start to finish, and what to check before sending a quote. It is written for people with no technical background.",
    )
    intro.paragraph_format.space_after = Pt(12)

    add_heading(doc, "1. What is the Quote Generator?", 1)
    add_body(
        doc,
        "The Quote Generator is an internal Jotform tool used to prepare a professional Enterprise quote for a customer or reseller. You enter the customer details, choose products and pricing, select any approved discounts and contract lengths, and review the finished quote on the screen.",
    )
    add_body(
        doc,
        "When everything is correct, the tool creates a PDF that can be saved and shared through the normal approval and sales process.",
    )

    add_heading(doc, "The aim", 2)
    add_bullet(doc, "Make quote preparation faster and more consistent.")
    add_bullet(doc, "Reduce manual calculations and copy-and-paste errors.")
    add_bullet(doc, "Show the customer or reseller a clear breakdown of products, quantities, term length, discounts, and total cost.")
    add_bullet(doc, "Produce a clean PDF that is ready for review and sharing.")

    add_heading(doc, "Who should use it?", 2)
    add_body(
        doc,
        "It is designed for sales team members and other authorized employees who prepare Jotform Enterprise quotes. You do not need technical knowledge. You only need the correct customer information, approved pricing choices, and any approved discount details.",
    )
    add_note(
        doc,
        "Important",
        "The tool helps prepare the quote, but it does not replace your company’s pricing, approval, legal, or tax rules. Confirm unusual pricing or discounts with the appropriate approver.",
    )

    add_heading(doc, "2. The process at a glance", 1)
    add_step(doc, "Open the tool.", "Enter the access password if the password screen appears.", 92)
    add_step(doc, "Choose who will receive the quote.", "Select Customer or Reseller.", 92)
    add_step(doc, "Enter the quote details.", "Add names, addresses, salesperson, Jotform entity, and currency.", 92)
    add_step(doc, "Add products and quantities.", "Use the product list, enter the quantity, and check the calculated amounts.", 92)
    add_step(doc, "Choose discounts and quote terms.", "Only select discounts that have been approved.", 92)
    add_step(doc, "Review the live preview.", "Read the quote as if you were the recipient.", 92)
    add_step(doc, "Download the PDF.", "Save it with the automatically prepared file name and complete any required internal review.", 92)

    add_note(
        doc,
        "Do not lose your work",
        "The Refresh button resets the form to a blank quote. Download and save the PDF before selecting Refresh or leaving the page.",
    )

    add_heading(doc, "3. Step-by-step instructions", 1)
    add_heading(doc, "Step 1 — Open the tool", 2)
    add_body(
        doc,
        "Open the Quote Generator link provided by your team. If you see a password page, enter the current shared or assigned password and select Continue. If the password is rejected, confirm that you are using the latest password.",
    )

    add_heading(doc, "Step 2 — Complete Quote Details", 2)
    add_body(doc, "Start in the Quote Details section on the left side of the screen.")
    add_bullet(doc, "Recipient type: Choose Customer for a direct quote. Choose Reseller when the quote is being sent through a reseller.")
    add_bullet(doc, "Quote date: This is filled in automatically using today’s date.")
    add_bullet(doc, "Contact and company details: Enter the recipient’s full name, company name, and complete address.")
    add_bullet(doc, "Reseller details: When Reseller is selected, enter both the reseller information and the end-customer information.")
    add_bullet(doc, "Salesperson name: Select the person preparing the quote. Their email and the quote number are filled in automatically.")
    add_bullet(doc, "Jotform entity: Select the correct Jotform legal entity. This updates the seller address and tax identification shown on the quote.")
    add_note(
        doc,
        "Check carefully",
        "The recipient type and Jotform entity affect what appears on the final document. If you are unsure which entity to use, ask before downloading the quote.",
    )

    add_heading(doc, "Step 3 — Add Quote Items", 2)
    add_body(doc, "In Quote Items, choose the currency first. Available currencies are USD, CAD, GBP, EUR, and AUD.")
    add_step(doc, "Select a product.", "Products are grouped as Base Packages, Add-Ons, and One-Time Fees.", 93)
    add_step(doc, "Enter the quantity.", "For example, enter 1 for one package or the required number of additional users.", 93)
    add_step(doc, "Check the unit price.", "The standard price appears automatically for the selected currency. Change it only when an approved custom amount is required.", 93)
    add_step(doc, "Check the total.", "The tool multiplies the unit price by the quantity and, for recurring items, by the selected term.", 93)
    add_step(doc, "Add more rows if needed.", "Select Add row for each additional product or fee. Use the arrows to change the order and the bin icon to remove a row.", 93)
    add_body(
        doc,
        "For Add another product or Professional Services, a custom product name and/or custom amount may be required. The PDF cannot be downloaded until the required information is entered.",
    )
    add_body(
        doc,
        "For a one-time fee, you may see a Waived checkbox. Select it only when the fee has been approved as waived. The line will show Waived and its amount will become zero.",
    )

    add_heading(doc, "Step 4 — Select discounts", 2)
    add_body(
        doc,
        "Use the Discount section to select an Education or Non-profit discount when the customer is eligible. Reseller discounts appear only when the recipient type is Reseller.",
    )
    add_note(
        doc,
        "Approval reminder",
        "Do not use a discount simply because it is available in the list. Confirm eligibility and approval first.",
    )

    add_heading(doc, "Step 5 — Select quote terms", 2)
    add_body(
        doc,
        "A quote term is the length of the agreement. You can select one or more term options so the recipient can compare them.",
    )
    add_bullet(doc, "12 months: no term discount.")
    add_bullet(doc, "2 years: 5% term discount.")
    add_bullet(doc, "3 years: 10% term discount.")
    add_bullet(doc, "5 years: 15% term discount.")
    add_bullet(doc, "Custom term: enter a clear label, number of months, and an optional discount percentage.")
    add_body(
        doc,
        "If more than one term is selected, the preview shows separate options labeled A, B, and so on. Review every option, not only the first one.",
    )

    add_heading(doc, "Step 6 — Add notes", 2)
    add_body(
        doc,
        "Use Add note for information that is specific to this quote, such as an agreed condition or clarification. Keep notes short, factual, and approved. The preview also contains standard notes such as quote validity and tax wording.",
    )

    add_heading(doc, "Step 7 — Review the preview", 2)
    add_body(
        doc,
        "The right side of the screen is the live preview. It updates as you make changes. Some text in the preview can also be edited directly, including headings, seller information, recipient details, labels, and standard notes.",
    )
    add_body(doc, "Before downloading, confirm all of the following:")
    add_bullet(doc, "Customer and reseller names are spelled correctly.")
    add_bullet(doc, "Addresses are complete and placed under the correct party.")
    add_bullet(doc, "Salesperson name, email, date, and quote number are correct.")
    add_bullet(doc, "The correct Jotform entity, address, and tax identifier are shown.")
    add_bullet(doc, "Currency, products, quantities, and unit prices are correct.")
    add_bullet(doc, "All discounts and waived fees are approved.")
    add_bullet(doc, "Every selected term shows the expected total.")
    add_bullet(doc, "Notes are accurate, necessary, and easy for the recipient to understand.")

    add_heading(doc, "Step 8 — Download the PDF", 2)
    add_body(
        doc,
        "Select Download PDF. Your browser’s print window opens. Choose Save as PDF, select the correct folder, and save the file. If the tool reports a missing custom product name or price, complete the highlighted field and try again.",
    )
    add_body(
        doc,
        "Open the saved PDF and check it one last time. Then follow your team’s normal approval and sending process.",
    )

    add_heading(doc, "4. How the totals work — in plain language", 1)
    add_body(
        doc,
        "The Quote Generator performs the arithmetic for you. Understanding the basic order helps you check whether the result makes sense.",
    )
    add_step(doc, "Calculate each line.", "Unit price × quantity × agreement length for recurring items. One-time fees are counted once.", 94)
    add_step(doc, "Separate recurring and one-time charges.", "This makes it clear which costs repeat and which happen only once.", 94)
    add_step(doc, "Apply eligible standard discounts.", "Education/non-profit discounts apply only to eligible products. Reseller discounts apply when a reseller quote is selected.", 94)
    add_step(doc, "Apply the term discount.", "The selected multi-year or custom term discount is applied to the applicable recurring amount.", 94)
    add_step(doc, "Add the remaining one-time charges.", "Waived one-time fees contribute zero.", 94)
    add_step(doc, "Show the final total.", "The preview displays the result for each selected term option.", 94)
    add_note(
        doc,
        "Simple sense-check",
        "A longer term normally increases the total contract value because it covers more years, even when a percentage discount is applied. Compare the term length and total together.",
    )

    add_heading(doc, "5. Worked example", 1)
    add_body(
        doc,
        "Imagine you are creating a direct customer quote in USD for one Jotform Enterprise Base Package and one Enterprise Onboarding fee.",
    )
    add_step(doc, "Choose Customer.", "Enter the customer’s contact, company, and address.", 95)
    add_step(doc, "Select the salesperson and Jotform entity.", "Check that the automatically filled details are correct.", 95)
    add_step(doc, "Choose USD.", "Add the base package with quantity 1 and the onboarding fee with quantity 1.", 95)
    add_step(doc, "Choose the term.", "Select 12 months for a one-year quote, or select more than one term to offer comparisons.", 95)
    add_step(doc, "Review the amounts.", "The package is treated as recurring; onboarding is treated as a one-time charge.", 95)
    add_step(doc, "Check and download.", "Read the preview, save the PDF, open it, and complete the normal approval process.", 95)

    add_heading(doc, "6. Common mistakes and how to avoid them", 1)
    add_heading(doc, "Wrong recipient type", 2)
    add_body(
        doc,
        "A reseller quote needs reseller details as well as end-customer details. Choose the recipient type before filling the rest of the form.",
    )
    add_heading(doc, "Wrong currency or entity", 2)
    add_body(
        doc,
        "Currency changes the prices shown. The Jotform entity changes the seller address and tax identifier. Check both near the beginning and again in the preview.",
    )
    add_heading(doc, "Unapproved discount or waived fee", 2)
    add_body(
        doc,
        "The presence of an option does not mean it is automatically allowed. Confirm approval before selecting it.",
    )
    add_heading(doc, "Missing custom information", 2)
    add_body(
        doc,
        "Custom products and certain services require a name or amount. Complete every field marked Required before downloading the PDF.",
    )
    add_heading(doc, "Accidentally clearing the quote", 2)
    add_body(
        doc,
        "Refresh returns the tool to a blank starting point. Save the PDF before using Refresh or closing the page.",
    )
    add_heading(doc, "Sending without opening the PDF", 2)
    add_body(
        doc,
        "Always open the saved PDF. Check page breaks, totals, names, addresses, and notes before it is reviewed or sent.",
    )

    add_heading(doc, "7. Final checklist", 1)
    checklist_items = [
        "Correct recipient type: Customer or Reseller",
        "Correct customer, company, reseller, and address details",
        "Correct salesperson, email, date, and quote number",
        "Correct Jotform entity and tax information",
        "Correct currency",
        "Correct products, quantities, and unit prices",
        "Custom product names and amounts completed",
        "Discounts and waived fees approved",
        "Correct term option or options",
        "Totals checked for reasonableness",
        "Notes reviewed",
        "PDF downloaded, opened, and checked",
        "Required internal approval completed before sending",
    ]
    for item in checklist_items:
        add_bullet(doc, f"☐ {item}")

    add_heading(doc, "8. Quick questions", 1)
    add_heading(doc, "Can I edit the preview?", 2)
    add_body(
        doc,
        "Yes. Many words and details in the preview can be edited directly. Use this carefully, because those changes appear in the final PDF.",
    )
    add_heading(doc, "Can I offer more than one term?", 2)
    add_body(
        doc,
        "Yes. Select multiple term options. The quote will display separate choices so the recipient can compare them.",
    )
    add_heading(doc, "Why can’t I download the PDF?", 2)
    add_body(
        doc,
        "A custom product name or custom amount is probably missing. Read the message above the preview and complete the highlighted field.",
    )
    add_heading(doc, "Does the tool save my draft?", 2)
    add_body(
        doc,
        "Treat the current screen as temporary. Download the PDF before refreshing, closing, or leaving the page.",
    )
    add_heading(doc, "What should I do if a price or rule looks wrong?", 2)
    add_body(
        doc,
        "Stop and confirm the current approved pricing or policy with the responsible team. Do not send a quote that you are uncertain about.",
    )

    add_heading(doc, "9. Getting help", 1)
    add_body(
        doc,
        "For access problems, incorrect product pricing, discount approval, entity selection, or unusual quote terms, contact the person or team responsible for the Quote Generator and Enterprise pricing in your organization. When asking for help, include a short description of what you were trying to do and a screenshot of the problem, but do not share sensitive customer information outside approved channels.",
    )

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_document()
    print(path.resolve())
